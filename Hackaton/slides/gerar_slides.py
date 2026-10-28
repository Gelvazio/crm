import re,zipfile
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

BG=RGBColor(0x05,0x03,0x0D); TXT=RGBColor(0xE6,0xF7,0xFF); SOFT=RGBColor(0x8F,0xA3,0xC7)
CY=RGBColor(0x22,0xE6,0xFF); MG=RGBColor(0xFF,0x2B,0xD6); VI=RGBColor(0x7B,0x2C,0xFF); GLASS=RGBColor(0x10,0x0A,0x24)
H="Orbitron"; D="Share Tech Mono"

slides=[
 ("HACKATHON DO FUTURO","AUTOTEC",[
  "Segurança para quem atravessa perto das escolas de Rio do Sul.",
  "Uma solução para scooters e bicicletas elétricas nas ciclofaixas.",
  "Nosso desafio: como podemos melhorar nossa cidade?"]),
 ("01 // O PROBLEMA","Scooters rápidas perto das escolas",[
  "Muitos autopropelidos circulam em Rio do Sul, inclusive perto das escolas.",
  "Crianças, jovens, pedestres e ciclistas correm risco ao atravessar.",
  "[X] de 9 alunos da nossa turma já viram isso acontecer."]),
 ("02 // A SOLUÇÃO","Como a Autotec funciona",[
  "A scooter tem GPS e conhece as zonas de risco da cidade.",
  "Perto da escola, ela desacelera sozinha e para antes da faixa de pedestres.",
  "Quem estava no limite de velocidade segue viagem normalmente."]),
 ("03 // FISCALIZAÇÃO","Quem passa do limite, para",[
  "Se a scooter estava acima do limite, a roda traseira trava.",
  "A polícia recebe um alerta e vai até o local desbloquear a scooter.",
  "As autopeças instalam o sistema, e vamos propor as zonas à Câmara de Vereadores."]),
 ("04 // POR QUE A AUTOTEC","Prevenir antes do primeiro acidente",[
  "Protege as crianças e também os pilotos, que muitas vezes são adolescentes.",
  "Ainda não há notícia de acidente aqui, e queremos que continue assim.",
  "Autotec: a tecnologia de Rio do Sul a favor de quem atravessa."]),
]

prs=Presentation(); prs.slide_width=Inches(10); prs.slide_height=Inches(5.625)
blank=prs.slide_layouts[6]

def alpha(shape,a):
    sf=shape.fill._xPr.find(qn('a:solidFill'))
    c=sf[0]; e=etree.SubElement(c,qn('a:alpha')); e.set('val',str(int(a*1000)))

def runs(slide,x,y,w,h,parts,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP):
    tb=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=tb.text_frame
    tf.word_wrap=True; tf.vertical_anchor=anchor
    for m in ("margin_left","margin_right","margin_top","margin_bottom"): setattr(tf,m,0)
    p=tf.paragraphs[0]; p.alignment=align
    for s,size,color,font,bold in parts:
        r=p.add_run(); r.text=s; r.font.size=Pt(size); r.font.color.rgb=color; r.font.name=font; r.font.bold=bold
    return tb

def glow(s,x,y,w,h,color,a):
    o=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(w),Inches(h))
    o.fill.solid(); o.fill.fore_color.rgb=color; alpha(o,a); o.line.fill.background(); o.shadow.inherit=False

def line(s,x,y,w,h,color,a=100):
    r=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    r.fill.solid(); r.fill.fore_color.rgb=color; r.line.fill.background(); r.shadow.inherit=False
    if a<100: alpha(r,a)

for i,(kicker,title,lines) in enumerate(slides):
    s=prs.slides.add_slide(blank)
    s.background.fill.solid(); s.background.fill.fore_color.rgb=BG
    glow(s,-1.5,-1.8,5.5,4.5,VI,22); glow(s,7.0,3.4,4.5,3.5,MG,14)
    for k in range(1,10): line(s,k*1.0,0,0.01,5.625,CY,6)
    for k in range(1,6): line(s,0,k*1.0,10,0.01,CY,6)
    # cabeçalho: Rio do Sul Mais Tec em destaque
    runs(s,0.5,0.3,6,0.3,[("RIO DO SUL ",13,TXT,H,True),("MAIS TEC",13,CY,H,True)])
    runs(s,6.0,0.32,3.5,0.3,[(kicker,10,MG,D,False)],align=PP_ALIGN.RIGHT)
    line(s,0.5,0.68,9.0,0.02,CY,70)
    if i==0:
        runs(s,0.5,1.0,9,1.0,[("AUTOTEC",54,TXT,H,True)])
        runs(s,0.5,1.95,9,0.4,[("Rio do Sul ",18,TXT,D,False),("Mais Tec",18,CY,D,True),("  ·  Hackathon do Futuro",18,SOFT,D,False)])
        top=2.55; h=0.62; gap=0.14; size=16
    else:
        runs(s,0.5,0.9,9,0.7,[(title,26,TXT,H,True)])
        top=1.75; h=0.88; gap=0.18; size=18
    for j,l in enumerate(lines):
        y=top+j*(h+gap)
        c=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(0.5),Inches(y),Inches(9.0),Inches(h))
        c.adjustments[0]=0.18; c.fill.solid(); c.fill.fore_color.rgb=GLASS; alpha(c,85)
        c.line.color.rgb=MG if (i==1 and j==2) else CY; c.line.width=Pt(1.25); c.shadow.inherit=False
        runs(s,0.75,y,0.6,h,[(f"0{j+1}",16,MG,H,True)],anchor=MSO_ANCHOR.MIDDLE)
        runs(s,1.4,y+0.05,7.9,h-0.1,[(l,size,TXT,D,False)],anchor=MSO_ANCHOR.MIDDLE)
    # rodapé com os três parceiros
    runs(s,0.5,5.2,7,0.25,[("RIO DO SUL MAIS TEC",10,CY,H,True),("  ·  ",10,SOFT,D,False),("SENAI",10,TXT,H,True),("  ·  ",10,SOFT,D,False),("<LAB365>",10,MG,H,True)])
    runs(s,8.5,5.2,1.0,0.25,[(f"{i+1}/5",10,SOFT,D,False)],align=PP_ALIGN.RIGHT)

used={id(x.slide_layout.part) for x in prs.slides}
for l in list(prs.slide_layouts):
    if id(l.part) not in used: prs.slide_layouts.remove(l)
prs.save("f0.pptx")
zin=zipfile.ZipFile("f0.pptx")
zout=zipfile.ZipFile("Autotec_futurista.pptx","w",zipfile.ZIP_DEFLATED,compresslevel=9)
for n in zin.namelist():
    if "printerSettings" in n or "thumbnail" in n: continue
    d=zin.read(n)
    if n=="[Content_Types].xml":
        for pat in (rb'<Default Extension="bin"[^>]*/>',rb'<Override PartName="/docProps/thumbnail[^>]*/>',rb'<Default Extension="jpeg"[^>]*/>'): d=re.sub(pat,b'',d)
    if n=="ppt/_rels/presentation.xml.rels": d=re.sub(rb'<Relationship [^>]*printerSettings[^>]*/>',b'',d)
    if n=="_rels/.rels": d=re.sub(rb'<Relationship [^>]*thumbnail[^>]*/>',b'',d)
    zout.writestr(n,d)
zout.close()
